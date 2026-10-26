"""Statystyka produkcji (dok. 6) — tabela `productions` w bazie M1.

Źródła wierszy:
  * EXCEL       — import z „2026 RAPORT_PRODUKCYJNY-STATYSTYKA” (rozwiązanie przeniesione z aplikacji M2-monika),
  * ROZLICZENIE — wiersz dopisany przy rozliczeniu koloru w module rozliczenia produkcji.
Jeden wiersz = jeden kolor jednego zamówienia wewnętrznego. Wiersz jest powiązany z klientem (AAA), produktem (BBBB),
formą (EEE), tworzywem SS (FFF) i barwnikiem BT (GGG) oraz ma kod szarży N-Z-S-DDMMRR.
"""
import csv
import io
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from flask import Response, abort, render_template, request, url_for
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from materials import auth

LOCK_ID = 6102033
STATS_XLSX = '2026 RAPORT_PRODUKCYJNY-STATYSTYKA (4).xlsx'
SKIP_SHEETS = re.compile(r'^(raport\b.*|arkusz\d+)$', re.I)
# Kolejność kolumn arkusza klienta (wiersz 2 = nagłówki, dane od wiersza 3) — jak w M2-monika.
COLUMNS = ['lp', 'product_name', 'start_date', 'end_date', 'internal_order', 'machine_no', 'qty', 'cycle_s',
           'avg_per_shift', 'cavities', 'real_cycle_s', 'scrap_pcs', 'net_g', 'shot_g', 'piece_gross_g', 'pcs_per_carton',
           'carton_kg', 'carton_type', 'pattern_no', 'material', 'colorant', 'dosage_pct', 'mounted_by', 'dismounted_by',
           'pattern_at_end', 'notes']
HEADERS = ['Lp.', 'Nazwa produktu', 'Data rozpoczęcia produkcji', 'Data zakończenia produkcji', 'Nr zamówienia wewn.',
           'Nr wtryskarki', 'Ilość łączna wyprodukowanych elementów', 'Czas cyklu [s] - na maszynie',
           'Średnia ilość produkowanych detali na 1 zmianę [8h]', 'Krotność formy',
           'Realny czas cyklu (uwzględniający brakowe cykle)', 'Szacunkowy poziom braków [szt.]', 'Waga netto 1 sztuki [g]',
           'Waga brutto - cały wtrysk z wlewkiem [g]', 'Waga 1 sztuki brutto [g]', 'Ilość sztuk w kartonie', 'Waga kartonu [kg]',
           'Rodzaj kartonu', 'Wzór produkcyjny nr', 'Symbol tworzywa', 'Symbol barwnika', 'Dozowanie barwnika [%]',
           'Kto zakładał formę', 'Kto demontował formę', 'Wzór z zakończenia produkcji', 'Uwagi']
DATES = {'start_date', 'end_date'}
NUMBERS = {'qty', 'cycle_s', 'avg_per_shift', 'cavities', 'real_cycle_s', 'scrap_pcs', 'net_g', 'shot_g', 'piece_gross_g',
           'pcs_per_carton', 'carton_kg', 'dosage_pct'}

# Forma wiersza: dla importu z Excela — aktualna forma produktu (poprawka w karcie produktu poprawia historię).
ROW_MOLD = "(CASE WHEN s.source = 'EXCEL' THEN pr.mold_id ELSE s.mold_id END)"
CODED_SQL = f'''
SELECT s.*, c.code AS client_code, c.short_name AS client_name, pr.code AS product_code, {ROW_MOLD} AS row_mold_id,
       mo.mold_no, ma.material_index, cb.material_index AS colorant_index, col.code AS color_code
FROM productions s
JOIN clients c ON c.id = s.client_id
LEFT JOIN products pr ON pr.id = s.product_id
LEFT JOIN molds mo ON mo.id = {ROW_MOLD}
LEFT JOIN materials ma ON ma.id = pr.material_id
LEFT JOIN materials cb ON cb.id = COALESCE(s.colorant_id, pr.colorant_id)
LEFT JOIN colors col ON col.id = pr.color_id'''


def text(value):
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    value = re.sub(r'\s+', ' ', str(value).replace(' ', ' ')).strip()
    return value or None


def number(value):
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    try:
        return Decimal(str(value).replace(',', '.').strip())
    except InvalidOperation:
        return None  # np. #DIV/0!, tekst


def as_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    match = re.fullmatch(r'(\d{1,2})[.,](\d{1,2})[.,](\d{4})', text(value) or '')
    if match:  # np. '14,09.2026'
        try:
            return date(int(match.group(3)), int(match.group(2)), int(match.group(1)))
        except ValueError:
            return None
    return None


def code3(index):
    return f'{int(index[2:]):03d}' if index and index[2:].isdigit() else '000'


def row_codes(row):
    """Kod produkcji AAA-BBBB-CCC i pełny kod AAA-BBBB-CCC-EEE-FFF-GGG dla wiersza statystyki."""
    if not row.get('product_code'):
        return None, None
    production = f'{row["client_code"]}-{row["product_code"]:04d}-{row.get("color_code") or "000"}'
    return production, f'{production}-{row.get("mold_no") or "000"}-{code3(row.get("material_index"))}-{code3(row.get("colorant_index"))}'


def init_stats(db, source_dir):
    with db() as connection, connection.cursor() as cursor:
        cursor.execute('SELECT pg_advisory_xact_lock(%s)', (LOCK_ID,))
        cursor.execute('''CREATE TABLE IF NOT EXISTS productions (
            id BIGSERIAL PRIMARY KEY, client_id BIGINT NOT NULL REFERENCES clients(id) ON DELETE RESTRICT,
            product_id BIGINT REFERENCES products(id) ON DELETE SET NULL, mold_id BIGINT REFERENCES molds(id) ON DELETE SET NULL,
            colorant_id BIGINT REFERENCES materials(id) ON DELETE SET NULL,
            source TEXT NOT NULL DEFAULT 'EXCEL', prod_color_id BIGINT UNIQUE, lp INTEGER, product_name TEXT NOT NULL,
            start_date DATE, end_date DATE, internal_order TEXT, machine_no TEXT, qty NUMERIC, cycle_s NUMERIC,
            avg_per_shift NUMERIC, cavities NUMERIC, real_cycle_s NUMERIC, scrap_pcs NUMERIC, cycles NUMERIC,
            net_g NUMERIC, shot_g NUMERIC, piece_gross_g NUMERIC, pcs_per_carton NUMERIC, carton_kg NUMERIC, carton_type TEXT,
            pattern_no TEXT, material TEXT, colorant TEXT, dosage_pct NUMERIC, mounted_by TEXT, dismounted_by TEXT,
            pattern_at_end TEXT, notes TEXT, batch_code TEXT, source_sheet TEXT, source_row INTEGER,
            created_at TIMESTAMPTZ NOT NULL DEFAULT NOW())''')
        cursor.execute('CREATE INDEX IF NOT EXISTS productions_client_idx ON productions(client_id, start_date)')
        cursor.execute('CREATE INDEX IF NOT EXISTS productions_product_idx ON productions(product_id)')
        cursor.execute('CREATE INDEX IF NOT EXISTS productions_order_idx ON productions(internal_order)')
        cursor.execute("SELECT COUNT(*) AS n FROM productions WHERE source = 'EXCEL'")
        path = Path(source_dir) / STATS_XLSX
        if cursor.fetchone()['n'] == 0 and path.exists():
            report = import_workbook(cursor, path)
            print('Import statystyki produkcji:', report, flush=True)


def import_workbook(cursor, path):
    cursor.execute('SELECT id, short_name FROM clients')
    clients = {r['short_name'].upper(): r['id'] for r in cursor.fetchall()}
    cursor.execute('SELECT id, client_id, UPPER(name) AS name FROM products')
    products = {(r['client_id'], re.sub(r'\s+', ' ', r['name']).strip()): r['id'] for r in cursor.fetchall()}
    workbook = load_workbook(path, read_only=True, data_only=True)
    total, unmatched, missing_clients = 0, 0, set()
    for ws in workbook.worksheets:
        sheet = text(ws.title)
        if SKIP_SHEETS.match(sheet):
            continue
        client_id = clients.get(sheet.upper())
        for row_no, row in enumerate(ws.iter_rows(min_row=3, max_col=len(COLUMNS), values_only=True), start=3):
            if not text(row[1]):
                continue
            if client_id is None:
                missing_clients.add(sheet)
                continue
            record = {'client_id': client_id, 'source_sheet': sheet, 'source_row': row_no}
            for column, value in zip(COLUMNS, row):
                if column in DATES:
                    record[column] = as_date(value)
                elif column in NUMBERS:
                    record[column] = number(value)
                else:
                    record[column] = text(value)
            lp = number(record['lp'])
            record['lp'] = int(lp) if lp is not None else None
            if record['dosage_pct'] is not None and record['dosage_pct'] < 1:
                record['dosage_pct'] = record['dosage_pct'] * 100  # arkusz trzyma 0,02 = 2%
            record['product_id'] = products.get((client_id, record['product_name'].upper()))
            unmatched += record['product_id'] is None
            fields = list(record)
            cursor.execute(f'INSERT INTO productions ({", ".join(fields)}) VALUES ({", ".join("%(" + f + ")s" for f in fields)})', record)
            total += 1
    workbook.close()
    assign_batch_codes(cursor)
    return {'wiersze': total, 'bez_produktu': unmatched, 'nieznani_klienci': sorted(missing_clients)}


def assign_batch_codes(cursor):
    """Kod szarży N-Z-S-DDMMRR dla historii. Szarża = całe zamówienie wewnętrzne na maszynie (instrukcja, słowniczek).
    Statystyka nie zapisuje zmiany startu, więc dla historii przyjęto Z = 1."""
    cursor.execute('''
        WITH runs AS (
            SELECT client_id, internal_order, machine_no, MIN(start_date) AS start
            FROM productions WHERE source = 'EXCEL' AND machine_no ~ '^[1-9]$' AND start_date IS NOT NULL
            GROUP BY client_id, internal_order, machine_no
        ), ranked AS (
            SELECT runs.*, ROW_NUMBER() OVER (PARTITION BY machine_no, start
                ORDER BY substring(internal_order from '^[0-9]+')::int NULLS LAST, client_id) AS s
            FROM runs
        )
        UPDATE productions p SET batch_code = r.machine_no || '-1-' || r.s || '-' || to_char(r.start, 'DDMMYY')
        FROM ranked r WHERE p.source = 'EXCEL' AND p.client_id = r.client_id AND p.machine_no = r.machine_no
          AND p.internal_order IS NOT DISTINCT FROM r.internal_order''')


# --- przeglądarka statystyki (jak w M2-monika) ---------------------------------------------------

TABLE = [
    ('batch_code', 'Kod szarży', 's.batch_code'), ('client', 'Klient', 'c.code'), ('product_code', 'BBBB', 'pr.code'),
    ('product_name', 'Produkt', 'lower(s.product_name)'), ('start_date', 'Start', 's.start_date'), ('end_date', 'Koniec', 's.end_date'),
    ('internal_order', 'Zam. wewn.', 's.internal_order'), ('machine_no', 'Masz.', 's.machine_no'), ('qty', 'Dobre szt.', 's.qty'),
    ('scrap_pcs', 'Braki szt.', 's.scrap_pcs'), ('cycle_s', 'Cykl s', 's.cycle_s'), ('avg_per_shift', 'Śr./zmianę', 's.avg_per_shift'),
    ('cavities', 'Krotn.', 's.cavities'), ('net_g', 'Netto g', 's.net_g'), ('shot_g', 'Wtrysk g', 's.shot_g'),
    ('mold_no', 'EEE', 'mo.mold_no'), ('material', 'Tworzywo', 'lower(s.material)'), ('colorant', 'Barwnik', 'lower(s.colorant)'),
    ('dosage_pct', 'Doz. %', 's.dosage_pct'), ('pattern_no', 'Wzór', 's.pattern_no'), ('source', 'Źródło', 's.source'),
]
SORTS = {k: v for k, _, v in TABLE}


def build_query(args):
    where, params = [], []
    clients = [c for c in args.getlist('klient') if c]
    if clients:
        where.append('c.code = ANY(%s)'); params.append(clients)
    for field, column in (('szarza', 's.batch_code'), ('q', 's.product_name'), ('tworzywo', 's.material'), ('barwnik', 's.colorant'),
                          ('zam', 's.internal_order'), ('uwagi', 's.notes'), ('wzor', 's.pattern_no')):
        value = args.get(field, '').strip()
        if value:
            where.append(f'{column} ILIKE %s'); params.append(f'%{value}%')
    for field, column in (('kod_produktu', "LPAD(pr.code::text, 4, '0')"), ('forma', 'mo.mold_no')):
        value = args.get(field, '').strip()
        if value:
            where.append(f'{column} = %s'); params.append(value.zfill(4 if field == 'kod_produktu' else 3))
    if args.get('maszyna', '').strip():
        where.append('s.machine_no = %s'); params.append(args['maszyna'].strip())
    if args.get('zrodlo') in ('EXCEL', 'ROZLICZENIE'):
        where.append('s.source = %s'); params.append(args['zrodlo'])
    for field, op in (('od', '>='), ('do', '<=')):
        value = args.get(field, '').strip()
        if value:
            where.append(f's.start_date {op} %s'); params.append(value)
    sort = args.get('sort', 'start_date')
    sort = sort if sort in SORTS else 'start_date'
    direction = 'ASC' if args.get('dir') == 'asc' else 'DESC'
    return ('WHERE ' + ' AND '.join(where)) if where else '', params, f'{SORTS[sort]} {direction} NULLS LAST, s.id', sort, direction.lower()


def register_stats_routes(app, db):
    def fetch(args, limit=None, offset=0):
        clause, params, order, sort, direction = build_query(args)
        with db() as connection, connection.cursor() as cursor:
            cursor.execute(f'''SELECT COUNT(*) AS n, COUNT(DISTINCT s.product_name) AS products, COUNT(DISTINCT c.id) AS clients,
                                      COALESCE(SUM(s.qty), 0) AS qty, COALESCE(SUM(s.scrap_pcs), 0) AS scrap,
                                      COUNT(DISTINCT s.internal_order) AS orders
                               FROM productions s JOIN clients c ON c.id = s.client_id LEFT JOIN products pr ON pr.id = s.product_id
                               LEFT JOIN molds mo ON mo.id = {ROW_MOLD} {clause}''', params)
            stats = cursor.fetchone()
            cursor.execute(f'{CODED_SQL} {clause} ORDER BY {order}' + (' LIMIT %s OFFSET %s' if limit else ''),
                           params + ([limit, offset] if limit else []))
            rows = cursor.fetchall()
            cursor.execute('''SELECT c.code, c.short_name, COUNT(s.id) AS n FROM clients c LEFT JOIN productions s ON s.client_id = c.id
                              GROUP BY c.id ORDER BY c.code''')
            clients = cursor.fetchall()
        for r in rows:
            r['production_code'], r['full_code'] = row_codes(r)
        return stats, rows, clients, sort, direction

    @app.get('/statystyka')
    @auth
    def stats_index():
        try:
            page = max(1, int(request.args.get('strona', 1)))
        except ValueError:
            page = 1
        stats, rows, clients, sort, direction = fetch(request.args, 100, (page - 1) * 100)
        pages = max(1, -(-stats['n'] // 100))
        return render_template('statistics.html', rows=rows, stats=stats, clients=clients, columns=TABLE, page=page, pages=pages,
                               sort=sort, direction=direction, selected=request.args.getlist('klient'), args=request.args)

    @app.template_global()
    def stats_url(**changes):
        args = request.args.to_dict(flat=False)
        for key, value in changes.items():
            if value is None:
                args.pop(key, None)
            else:
                args[key] = [value]
        return url_for('stats_index', **args)

    @app.get('/statystyka/eksport.csv')
    @auth
    def stats_csv():
        _, rows, *_ = fetch(request.args)
        if not rows:
            abort(404)
        out = io.StringIO()
        writer = csv.writer(out, delimiter=';')
        keys = ['client_code', 'client_name', 'batch_code', 'production_code', 'full_code'] + COLUMNS + ['cycles', 'source']
        writer.writerow(keys)
        for row in rows:
            writer.writerow(['' if row.get(k) is None else row.get(k) for k in keys])
        return Response('﻿' + out.getvalue(), mimetype='text/csv',
                        headers={'Content-Disposition': 'attachment; filename=statystyka_produkcji.csv'})

    @app.get('/statystyka/eksport.xlsx')
    @auth
    def stats_xlsx():
        """Eksport w układzie arkusza „STATYSTYKI PRODUKCJI” — jedna zakładka na klienta."""
        _, rows, *_ = fetch(request.args)
        workbook = Workbook()
        workbook.remove(workbook.active)
        head_fill = PatternFill('solid', fgColor='D9E1F2')
        by_client = {}
        for r in sorted(rows, key=lambda r: (r['client_name'], r['start_date'] or date.min, r['id'])):
            by_client.setdefault(r['client_name'], []).append(r)
        for client, items in by_client.items():
            ws = workbook.create_sheet(client[:31])
            ws.append(['STATYSTYKI PRODUKCJI', None, None, None, None, None, None, client])
            ws.append(HEADERS + ['Kod szarży', 'Kod produkcji', 'Pełny kod'])
            for cell in ws[2]:
                cell.font, cell.fill = Font(bold=True), head_fill
                cell.alignment = Alignment(wrap_text=True, vertical='top')
            for n, r in enumerate(items, start=1):
                values = [r.get(c) for c in COLUMNS]
                values[0] = f'{n}.'
                ws.append([float(v) if isinstance(v, Decimal) else v for v in values] + [r['batch_code'], r['production_code'], r['full_code']])
            ws.freeze_panes = 'C3'
        if not by_client:
            workbook.create_sheet('brak danych')
        buffer = io.BytesIO()
        workbook.save(buffer)
        return Response(buffer.getvalue(), mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                        headers={'Content-Disposition': 'attachment; filename=STATYSTYKA_PRODUKCJI.xlsx'})
