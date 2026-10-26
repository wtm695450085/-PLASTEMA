import hmac
import os
import secrets
from functools import wraps
from pathlib import Path

import psycopg
from flask import Flask, abort, flash, redirect, render_template, request, session, url_for
from openpyxl import load_workbook
from psycopg.errors import UniqueViolation
from psycopg.rows import dict_row
from werkzeug.security import check_password_hash

import shell
from changelog import CHANGELOG

DATABASE_URL = os.environ["DATABASE_URL"]
SOURCE_XLSX = Path(os.getenv("SOURCE_XLSX", "/app/source/BAZA DOSTAWCÓW 1.xlsx"))
SUPPLIER_FIELDS = ["ordinal", "name", "short_name", "supplier_no", "material_type", "address",
                   "postal_code", "city", "country", "nip", "phone", "website", "payment_terms", "notes"]
CONTACT_FIELDS = ["contact_name", "phone", "email", "position", "notes"]

app = Flask(__name__)
app.config.update(
    SECRET_KEY=os.environ["SECRET_KEY"], SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
    SESSION_COOKIE_SECURE=os.getenv("COOKIE_SECURE", "true").lower() == "true",
    PERMANENT_SESSION_LIFETIME=28800,
)


def db():
    return psycopg.connect(DATABASE_URL, row_factory=dict_row)


def clean(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).replace("\u00a0", " ").strip()


def seed_from_workbook(cursor):
    workbook = load_workbook(SOURCE_XLSX, read_only=True, data_only=True)
    sheet = workbook["DOSTAWCY TWORZYW"]
    for row in sheet.iter_rows(min_row=5, max_col=15, values_only=True):
        values = [clean(value) for value in row[1:15]]
        if not values[1] or not values[3]:
            continue
        try:
            values[0] = int(values[0]) if values[0] else None
        except ValueError:
            values[0] = None
        values[3] = values[3] or None
        placeholders = ",".join("%s" for _ in SUPPLIER_FIELDS)
        cursor.execute(
            f"INSERT INTO suppliers ({','.join(SUPPLIER_FIELDS)}) VALUES ({placeholders}) "
            "ON CONFLICT (supplier_no) DO NOTHING", values,
        )

    sheet = workbook["DOSTAWCY BARWNIKÓW"]
    for row in sheet.iter_rows(min_row=5, max_col=15, values_only=True):
        values = [clean(value) for value in row[1:15]]
        if not values[1] or not values[3]:
            continue
        try:
            values[0] = int(values[0]) if values[0] else None
        except ValueError:
            values[0] = None
        values[3] = values[3] or None
        placeholders = ",".join("%s" for _ in SUPPLIER_FIELDS)
        sql = "INSERT INTO suppliers (" + ",".join(SUPPLIER_FIELDS) + ") VALUES (" + placeholders + ") ON CONFLICT (supplier_no) DO NOTHING"
        cursor.execute(sql, values)

    sheet = workbook["DOSTAWCY TWORZ_OSOBY KONTAKT "]
    for row in sheet.iter_rows(min_row=5, max_col=8, values_only=True):
        ordinal, short_name, supplier_no, contact_name, phone, email, notes = [clean(value) for value in row[1:8]]
        if not supplier_no:
            continue
        cursor.execute("SELECT id FROM suppliers WHERE supplier_no = %s", (supplier_no,))
        supplier = cursor.fetchone()
        if supplier is None:
            try:
                ordinal_value = int(ordinal) if ordinal else None
            except ValueError:
                ordinal_value = None
            cursor.execute(
                "INSERT INTO suppliers (ordinal,name,short_name,supplier_no,material_type) "
                "VALUES (%s,%s,%s,%s,'TWORZYWO') RETURNING id",
                (ordinal_value, short_name or supplier_no, short_name, supplier_no),
            )
            supplier = cursor.fetchone()
        if any((contact_name, phone, email, notes)):
            cursor.execute(
                "INSERT INTO contacts (supplier_id,contact_name,phone,email,position,notes) VALUES (%s,%s,%s,%s,%s,%s)",
                (supplier["id"], contact_name, phone, email, "", notes),
            )
    sheet = workbook["DOSTAWCY BARW._OSOBY KONT. "]
    for row in sheet.iter_rows(min_row=5, max_col=11, values_only=True):
        _, supplier_name, short_name, supplier_no, _, contact_name, phone, email, position, notes = [clean(value) for value in row[1:11]]
        if not supplier_no:
            continue
        cursor.execute("SELECT id FROM suppliers WHERE supplier_no = %s", (supplier_no,))
        supplier = cursor.fetchone()
        if supplier is None:
            cursor.execute("INSERT INTO suppliers (name,short_name,supplier_no,material_type) VALUES (%s,%s,%s,%s) RETURNING id", (supplier_name or supplier_no, short_name, supplier_no, "BARWNIK"))
            supplier = cursor.fetchone()
        if any((contact_name, phone, email, position, notes)):
            cursor.execute("INSERT INTO contacts (supplier_id,contact_name,phone,email,position,notes) VALUES (%s,%s,%s,%s,%s,%s)", (supplier["id"], contact_name, phone, email, position, notes))

    workbook.close()


def init_db():
    with db() as connection, connection.cursor() as cursor:
        cursor.execute("SELECT pg_advisory_xact_lock(6102026)")
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS suppliers (
                id BIGSERIAL PRIMARY KEY, ordinal INTEGER, name TEXT NOT NULL, short_name TEXT,
                supplier_no TEXT UNIQUE, material_type TEXT, address TEXT, postal_code TEXT,
                city TEXT, country TEXT, nip TEXT, phone TEXT, website TEXT, payment_terms TEXT,
                notes TEXT, created_at TIMESTAMPTZ NOT NULL DEFAULT NOW(),
                updated_at TIMESTAMPTZ NOT NULL DEFAULT NOW()
            )
        """)
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS contacts (
                id BIGSERIAL PRIMARY KEY,
                supplier_id BIGINT NOT NULL REFERENCES suppliers(id) ON DELETE CASCADE,
                contact_name TEXT, phone TEXT, email TEXT, position TEXT, notes TEXT,
                created_at TIMESTAMPTZ NOT NULL DEFAULT NOW(),
                updated_at TIMESTAMPTZ NOT NULL DEFAULT NOW(),
                CONSTRAINT contacts_contact_required CHECK (NULLIF(contact_name,'') IS NOT NULL OR NULLIF(phone,'') IS NOT NULL
                       OR NULLIF(email,'') IS NOT NULL OR NULLIF(position,'') IS NOT NULL OR NULLIF(notes,'') IS NOT NULL)
            )
        """)
        cursor.execute("ALTER TABLE contacts ADD COLUMN IF NOT EXISTS position TEXT")
        cursor.execute("ALTER TABLE contacts DROP CONSTRAINT IF EXISTS contacts_check")
        cursor.execute("ALTER TABLE contacts DROP CONSTRAINT IF EXISTS contacts_contact_required")
        cursor.execute("ALTER TABLE contacts ADD CONSTRAINT contacts_contact_required CHECK (NULLIF(contact_name, CHR(39)||CHR(39)) IS NOT NULL OR NULLIF(phone, CHR(39)||CHR(39)) IS NOT NULL OR NULLIF(email, CHR(39)||CHR(39)) IS NOT NULL OR NULLIF(position, CHR(39)||CHR(39)) IS NOT NULL OR NULLIF(notes, CHR(39)||CHR(39)) IS NOT NULL)")
        cursor.execute("CREATE INDEX IF NOT EXISTS contacts_supplier_id_idx ON contacts(supplier_id)")
        cursor.execute("SELECT COUNT(*) AS count FROM suppliers")
        if cursor.fetchone()["count"] == 0 and SOURCE_XLSX.exists():
            seed_from_workbook(cursor)


def login_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        if not session.get("authenticated"):
            return redirect(url_for("login"))
        return view(*args, **kwargs)
    return wrapped


def csrf_token():
    if "csrf_token" not in session:
        session["csrf_token"] = secrets.token_urlsafe(32)
    return session["csrf_token"]


def require_csrf():
    if not hmac.compare_digest(request.form.get("csrf_token", ""), session.get("csrf_token", "")):
        abort(400, "Nieprawidłowy token formularza.")


app.jinja_env.globals["csrf_token"] = csrf_token
app.jinja_env.globals["icon"] = shell.icon
app.jinja_env.globals["logo"] = shell.logo


def attention_items():
    """Liczby z całej aplikacji, które wymagają reakcji (pulpit i kropka przy dzwonku)."""
    queries = [
        ("SELECT COUNT(*) AS n FROM orders WHERE status='NOWE'", "Nowe zamówienia czekają na przyjęcie", "Sprawdź formę, surowiec i maszynę, potem przyjmij do realizacji", "orders_index", {"status": "NOWE"}, "warn"),
        ("SELECT COUNT(*) AS n FROM orders WHERE status='PRZYJĘTE'", "Przyjęte zamówienia do przekazania do produkcji", "Przycisk „Przekaż do produkcji” na karcie zamówienia", "orders_index", {"status": "PRZYJĘTE"}, "info"),
        ("SELECT COUNT(*) AS n FROM prod_colors WHERE status='W PRODUKCJI'", "Kolory w produkcji do rozliczenia", "Raporty zmianowe i rozliczenie koloru", "settlement_index", {}, "info"),
        ("SELECT COUNT(*) AS n FROM fg_docs WHERE status='OTWARTY'", "Otwarte dokumenty magazynu wyrobów", "WZ, ZW lub INW czekają na realizację", "fg_docs", {}, "warn"),
        ("SELECT COUNT(*) AS n FROM receipts WHERE status='OTWARTY'", "Otwarte przyjęcia towaru", "Dostawy surowców czekające na zatwierdzenie", "receipts_index", {}, "warn"),
        ("SELECT COUNT(*) AS n FROM machine_docs WHERE status='OTWARTY'", "Otwarte dokumenty maszyn", "Wydania i zwroty surowca czekają na magazyniera", "machine_docs_index", {}, "warn"),
        ("SELECT COUNT(*) AS n FROM inventories WHERE status='OTWARTA'", "Otwarte inwentaryzacje", "Arkusz spisu nie został jeszcze zatwierdzony", "inventory_index", {}, "warn"),
    ]
    items = []
    with db() as connection, connection.cursor() as cursor:
        for sql, title, hint, endpoint, params, tone in queries:
            try:
                cursor.execute(sql)
                count = cursor.fetchone()["n"]
            except psycopg.Error:
                connection.rollback()
                continue
            if count:
                items.append({"count": count, "title": title, "hint": hint, "href": url_for(endpoint, **params), "tone": tone})
    return items


@app.context_processor
def inject_shell():
    if not session.get("authenticated") or request.endpoint in (None, "static"):
        return {}
    nav_top, nav_sections, nav_active, nav_crumbs = shell.build_nav(request, url_for)
    return {"nav_top": nav_top, "nav_sections": nav_sections, "nav_active": nav_active, "nav_crumbs": nav_crumbs,
            "ui_user": shell.user_info(), "attention_total": session.get("attention_total", 0)}


@app.get("/pulpit")
@login_required
def dashboard():
    items = attention_items()
    session["attention_total"] = sum(1 for i in items if i["tone"] == "warn")
    with db() as connection, connection.cursor() as cursor:
        cursor.execute("""SELECT
            (SELECT COUNT(*) FROM orders WHERE status IN ('NOWE','PRZYJĘTE','W PRODUKCJI')) AS open_orders,
            (SELECT COUNT(*) FROM prod_orders WHERE status NOT IN ('ZAKOŃCZONE')) AS open_prod,
            (SELECT COALESCE(SUM(qty),0) FROM fg_moves) AS fg_qty,
            (SELECT COUNT(*) FROM clients) AS clients""")
        kpi = cursor.fetchone()
    return render_template("dashboard.html", items=items, kpi=kpi, changelog=CHANGELOG)


@app.after_request
def security_headers(response):
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["X-Frame-Options"] = "DENY"
    response.headers["Referrer-Policy"] = "same-origin"
    response.headers["Content-Security-Policy"] = (
        "default-src 'self'; style-src 'self'; script-src 'self'; img-src 'self' data:; "
        "font-src 'self'; frame-ancestors 'none'; base-uri 'self'; form-action 'self'"
    )
    return response


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        require_csrf()
        username_ok = hmac.compare_digest(request.form.get("username", ""), os.getenv("APP_USERNAME", "admin"))
        password_ok = check_password_hash(os.environ["APP_PASSWORD_HASH"], request.form.get("password", ""))
        if username_ok and password_ok:
            session.clear()
            session["authenticated"] = True
            session.permanent = True
            return redirect(url_for("dashboard"))
        flash("Nieprawidłowy login lub hasło.", "error")
    return render_template("login.html")


@app.post("/logout")
@login_required
def logout():
    require_csrf()
    session.clear()
    return redirect(url_for("login"))


@app.get("/")
@login_required
def index():
    query, material = request.args.get("q", "").strip(), request.args.get("material", "").strip()
    sql = "SELECT s.*,COUNT(c.id) AS contact_count FROM suppliers s LEFT JOIN contacts c ON c.supplier_id=s.id WHERE TRUE"
    params = []
    if query:
        term = f"%{query}%"
        sql += " AND (s.name ILIKE %s OR s.short_name ILIKE %s OR s.supplier_no ILIKE %s OR s.city ILIKE %s OR s.nip ILIKE %s)"
        params.extend([term] * 5)
    if material:
        sql += " AND s.material_type=%s"
        params.append(material)
    sql += " GROUP BY s.id ORDER BY CASE s.material_type WHEN 'TWORZYWO' THEN 0 WHEN 'BARWNIK' THEN 1 ELSE 2 END,s.ordinal,s.name"
    with db() as connection, connection.cursor() as cursor:
        cursor.execute(sql, params)
        suppliers = cursor.fetchall()
        cursor.execute("SELECT DISTINCT material_type FROM suppliers WHERE material_type<>'' ORDER BY material_type")
        materials = [row["material_type"] for row in cursor.fetchall()]
        cursor.execute("SELECT COUNT(*) AS count FROM suppliers")
        total = cursor.fetchone()["count"]
        cursor.execute("SELECT COUNT(*) AS count FROM contacts")
        contacts_total = cursor.fetchone()["count"]
    return render_template("index.html", suppliers=suppliers, materials=materials, total=total,
                           contacts_total=contacts_total, query=query, material=material)


def supplier_form_values():
    values = {field: request.form.get(field, "").strip() for field in SUPPLIER_FIELDS}
    if not values["name"]:
        raise ValueError("Nazwa dostawcy jest wymagana.")
    try:
        values["ordinal"] = int(values["ordinal"]) if values["ordinal"] else None
    except ValueError as error:
        raise ValueError("Pole Lp. musi być liczbą.") from error
    values["supplier_no"] = values["supplier_no"] or None
    return values


def contact_form_values():
    values = {field: request.form.get(field, "").strip() for field in CONTACT_FIELDS}
    if not any(values.values()):
        raise ValueError("Wpisz co najmniej jedną informację o kontakcie.")
    return values


@app.route("/dostawcy/nowy", methods=["GET", "POST"])
@login_required
def create_supplier():
    if request.method == "POST":
        require_csrf()
        try:
            values = supplier_form_values()
            with db() as connection, connection.cursor() as cursor:
                placeholders = ",".join("%s" for _ in SUPPLIER_FIELDS)
                cursor.execute(f"INSERT INTO suppliers ({','.join(SUPPLIER_FIELDS)}) VALUES ({placeholders}) RETURNING id",
                               [values[field] for field in SUPPLIER_FIELDS])
                supplier_id = cursor.fetchone()["id"]
            flash("Dostawca został dodany. Możesz teraz dodać osoby kontaktowe.", "success")
            return redirect(url_for("supplier_contacts", supplier_id=supplier_id))
        except ValueError as error:
            flash(str(error), "error")
        except UniqueViolation:
            flash("Numer dostawcy musi być unikalny.", "error")
    return render_template("form.html", supplier=request.form or {}, contacts=[], title="Nowy dostawca")


@app.route("/dostawcy/<int:supplier_id>/edytuj", methods=["GET", "POST"])
@login_required
def edit_supplier(supplier_id):
    with db() as connection, connection.cursor() as cursor:
        cursor.execute("SELECT * FROM suppliers WHERE id=%s", (supplier_id,))
        supplier = cursor.fetchone()
        cursor.execute("SELECT * FROM contacts WHERE supplier_id=%s ORDER BY id", (supplier_id,))
        contacts = cursor.fetchall()
    if supplier is None:
        abort(404)
    if request.method == "POST":
        require_csrf()
        try:
            values = supplier_form_values()
            assignments = ",".join(f"{field}=%s" for field in SUPPLIER_FIELDS)
            with db() as connection, connection.cursor() as cursor:
                cursor.execute(f"UPDATE suppliers SET {assignments},updated_at=NOW() WHERE id=%s",
                               [values[field] for field in SUPPLIER_FIELDS] + [supplier_id])
            flash("Dane dostawcy zostały zapisane.", "success")
            return redirect(url_for("edit_supplier", supplier_id=supplier_id))
        except ValueError as error:
            flash(str(error), "error")
            supplier = request.form
        except UniqueViolation:
            flash("Numer dostawcy musi być unikalny.", "error")
            supplier = request.form
    return render_template("form.html", supplier=supplier, contacts=contacts, title="Edycja dostawcy")


@app.post("/dostawcy/<int:supplier_id>/usun")
@login_required
def delete_supplier(supplier_id):
    require_csrf()
    with db() as connection, connection.cursor() as cursor:
        cursor.execute("DELETE FROM suppliers WHERE id=%s", (supplier_id,))
        deleted = cursor.rowcount
    if deleted:
        flash("Dostawca wraz z powiązanymi kontaktami został usunięty.", "success")
    return redirect(url_for("index"))


@app.get("/dostawcy/<int:supplier_id>/kontakty")
@login_required
def supplier_contacts(supplier_id):
    with db() as connection, connection.cursor() as cursor:
        cursor.execute("SELECT id,name,short_name,supplier_no,material_type,phone FROM suppliers WHERE id=%s", (supplier_id,))
        supplier = cursor.fetchone()
        if supplier is None:
            abort(404)
        cursor.execute("SELECT * FROM contacts WHERE supplier_id=%s ORDER BY id", (supplier_id,))
        contacts = cursor.fetchall()
    return render_template("supplier_contacts.html", supplier=supplier, contacts=contacts)


@app.route("/dostawcy/<int:supplier_id>/kontakty/nowy", methods=["GET", "POST"])
@login_required
def create_contact(supplier_id):
    with db() as connection, connection.cursor() as cursor:
        cursor.execute("SELECT id,name,supplier_no FROM suppliers WHERE id=%s", (supplier_id,))
        supplier = cursor.fetchone()
    if supplier is None:
        abort(404)
    if request.method == "POST":
        require_csrf()
        try:
            values = contact_form_values()
            with db() as connection, connection.cursor() as cursor:
                cursor.execute("INSERT INTO contacts (supplier_id,contact_name,phone,email,position,notes) VALUES (%s,%s,%s,%s,%s,%s)",
                               (supplier_id, *[values[field] for field in CONTACT_FIELDS]))
            flash("Osoba kontaktowa została dodana.", "success")
            return redirect(url_for("supplier_contacts", supplier_id=supplier_id))
        except ValueError as error:
            flash(str(error), "error")
    return render_template("contact_form.html", contact=request.form or {}, supplier=supplier, title="Nowa osoba kontaktowa")


@app.route("/kontakty/<int:contact_id>/edytuj", methods=["GET", "POST"])
@login_required
def edit_contact(contact_id):
    with db() as connection, connection.cursor() as cursor:
        cursor.execute("SELECT c.*,s.name AS supplier_name,s.supplier_no FROM contacts c JOIN suppliers s ON s.id=c.supplier_id WHERE c.id=%s", (contact_id,))
        contact = cursor.fetchone()
    if contact is None:
        abort(404)
    supplier = {"id": contact["supplier_id"], "name": contact["supplier_name"], "supplier_no": contact["supplier_no"]}
    if request.method == "POST":
        require_csrf()
        try:
            values = contact_form_values()
            with db() as connection, connection.cursor() as cursor:
                cursor.execute("UPDATE contacts SET contact_name=%s,phone=%s,email=%s,position=%s,notes=%s,updated_at=NOW() WHERE id=%s",
                               (*[values[field] for field in CONTACT_FIELDS], contact_id))
            flash("Dane osoby kontaktowej zostały zapisane.", "success")
            return redirect(url_for("supplier_contacts", supplier_id=contact["supplier_id"]))
        except ValueError as error:
            flash(str(error), "error")
            contact = request.form
    return render_template("contact_form.html", contact=contact, supplier=supplier, title="Edycja osoby kontaktowej")


@app.post("/kontakty/<int:contact_id>/usun")
@login_required
def delete_contact(contact_id):
    require_csrf()
    with db() as connection, connection.cursor() as cursor:
        cursor.execute("DELETE FROM contacts WHERE id=%s RETURNING supplier_id", (contact_id,))
        deleted = cursor.fetchone()
    if deleted is None:
        abort(404)
    flash("Osoba kontaktowa została usunięta.", "success")
    return redirect(url_for("supplier_contacts", supplier_id=deleted["supplier_id"]))


@app.get("/health")
def health():
    with db() as connection, connection.cursor() as cursor:
        cursor.execute("SELECT 1")
    return {"status": "ok", "database": "postgresql"}


init_db()

from materials import init_materials, register_material_routes
init_materials(db, SOURCE_XLSX)
register_material_routes(app, db)

from inventory import init_inventory, register_inventory_routes
init_inventory(db)
register_inventory_routes(app, db)

from machines import init_machines, register_machine_routes
init_machines(db)
register_machine_routes(app, db)

from receipts import init_receipts, register_receipt_routes
init_receipts(db)
register_receipt_routes(app, db)

from orders import init_orders, register_order_routes
init_orders(db, os.getenv("ORDERS_SOURCE_DIR", "/app/source/ikony"))
register_order_routes(app, db)

from molds import init_molds, register_mold_routes
init_molds(db)
register_mold_routes(app, db)

from prod_stats import init_stats, register_stats_routes
init_stats(db, os.getenv("ORDERS_SOURCE_DIR", "/app/source/ikony"))
register_stats_routes(app, db)

from settlement import init_settlement, register_settlement_routes
init_settlement(db)
register_settlement_routes(app, db)

from fg import init_fg, register_fg_routes
init_fg(db)
register_fg_routes(app, db)
