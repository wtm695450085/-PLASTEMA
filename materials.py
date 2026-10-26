import hmac
import secrets
from decimal import Decimal, InvalidOperation
from functools import wraps

from flask import abort, flash, redirect, render_template, request, session, url_for
from openpyxl import load_workbook
from psycopg.errors import ForeignKeyViolation

MATERIAL_FIELDS = ['material_index','ordinal','supplier_id','supplier_no','supplier_name','supplier_short_name',
                   'material_type','kind','origin_code','symbol','color','internal_name','unit_price',
                   'currency','price_pln','unit','lead_time_days','notes']
FORM_FIELDS = ['ordinal','supplier_no','supplier_name','supplier_short_name','material_type','kind','origin_code',
               'symbol','color','internal_name','unit_price','currency','price_pln','unit','lead_time_days','notes']


def clean(value):
    if value is None: return ''
    if isinstance(value, float) and value.is_integer(): value = int(value)
    return str(value).replace('\u00a0',' ').strip()


def as_decimal(value):
    value = clean(value).replace(',', '.')
    if not value: return None
    try: return Decimal(value)
    except InvalidOperation as error: raise ValueError('Cena musi być liczbą.') from error


def as_int(value, label):
    value = clean(value)
    if not value: return None
    try: return int(value)
    except ValueError as error: raise ValueError(f'{label} musi być liczbą całkowitą.') from error


def index_prefix(material_type):
    return 'BT' if 'BARWNIK' in clean(material_type).upper() else 'SS'


def auth(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        if not session.get('authenticated'): return redirect(url_for('login'))
        return view(*args, **kwargs)
    return wrapped


def csrf():
    if 'csrf_token' not in session: session['csrf_token'] = secrets.token_urlsafe(32)
    return session['csrf_token']


def require_csrf():
    if not hmac.compare_digest(request.form.get('csrf_token',''), session.get('csrf_token','')): abort(400)


def next_index(cursor, prefix, exclude_id=None):
    if exclude_id is None:
        cursor.execute("SELECT COALESCE(MAX(CAST(SUBSTRING(material_index,3) AS INTEGER)),0)+1 AS next FROM materials WHERE material_index LIKE %s", (prefix+'%',))
    else:
        cursor.execute("SELECT COALESCE(MAX(CAST(SUBSTRING(material_index,3) AS INTEGER)),0)+1 AS next FROM materials WHERE material_index LIKE %s AND id<>%s", (prefix+'%', exclude_id))
    return f'{prefix}{cursor.fetchone()["next"]:03d}'


def ensure_barwnik_materials(cursor):
    return


def init_materials(db, source_xlsx):
    with db() as connection, connection.cursor() as cursor:
        cursor.execute('SELECT pg_advisory_xact_lock(6102027)')
        cursor.execute('''CREATE TABLE IF NOT EXISTS materials (
            id BIGSERIAL PRIMARY KEY, material_index TEXT NOT NULL UNIQUE, ordinal INTEGER,
            supplier_id BIGINT REFERENCES suppliers(id) ON DELETE SET NULL, supplier_no TEXT,
            supplier_name TEXT, supplier_short_name TEXT, material_type TEXT, kind TEXT,
            origin_code INTEGER, symbol TEXT, color TEXT, internal_name TEXT,
            unit_price NUMERIC(12,4), currency TEXT, price_pln NUMERIC(12,4), unit TEXT,
            lead_time_days INTEGER, notes TEXT, created_at TIMESTAMPTZ NOT NULL DEFAULT NOW(),
            updated_at TIMESTAMPTZ NOT NULL DEFAULT NOW())''')
        cursor.execute('CREATE INDEX IF NOT EXISTS materials_supplier_idx ON materials(supplier_id)')
        cursor.execute('CREATE INDEX IF NOT EXISTS materials_search_idx ON materials(kind,color,symbol)')
        cursor.execute("SELECT id,material_index,material_type FROM materials ORDER BY id")
        old_rows = cursor.fetchall()
        for row in old_rows:
            old = row['material_index'] or ''
            material_type = 'SUROWIEC SYPKI' if clean(row['material_type']).upper() == 'TWORZYWO' else row['material_type']
            new_index = old
            if old.startswith('TS'):
                new_index = 'SS' + old[2:]
            if material_type != row['material_type'] or new_index != old:
                cursor.execute('UPDATE materials SET material_type=%s,material_index=%s,updated_at=NOW() WHERE id=%s', (material_type,new_index,row['id']))
        cursor.execute('SELECT COUNT(*) AS count FROM materials')
        if cursor.fetchone()["count"] or not source_xlsx.exists():
            ensure_barwnik_materials(cursor)
            return
        wb = load_workbook(source_xlsx, read_only=True, data_only=True)
        ws = wb['WYKAZ TWORZYW']; idx = 1
        for row in ws.iter_rows(min_row=5, max_col=17, values_only=True):
            vals = [clean(v) for v in row[1:17]]
            if not vals[1] or not vals[4]: continue
            cursor.execute('SELECT id FROM suppliers WHERE supplier_no=%s',(vals[3],)); supplier=cursor.fetchone()
            cursor.execute('''INSERT INTO materials
                (material_index,ordinal,supplier_id,supplier_no,supplier_name,supplier_short_name,material_type,kind,origin_code,symbol,color,internal_name,unit_price,currency,price_pln,unit,lead_time_days,notes)
                VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)''',
                (f'SS{idx:03d}',as_int(vals[0],'Lp.'),supplier['id'] if supplier else None,vals[3],vals[1],vals[2],'SUROWIEC SYPKI',vals[5],as_int(vals[6],'Kod oryginał/regranulat'),vals[7],vals[8],vals[9],as_decimal(vals[10]),vals[11],as_decimal(vals[12]),vals[13],as_int(vals[14],'Czas dostawy'),vals[15]))
            idx += 1
        wb.close()
        ensure_barwnik_materials(cursor)


def register_material_routes(app, db):
    app.jinja_env.globals['material_csrf_token'] = csrf

    def form_values():
        values = {f: request.form.get(f,'').strip() for f in FORM_FIELDS}
        values['material_type'] = values['material_type'] or 'SUROWIEC SYPKI'
        if values['material_type'] not in ('SUROWIEC SYPKI','BARWNIK TWORZYWA'):
            raise ValueError('Wybierz kategorię SS — surowiec sypki albo BT — barwnik tworzywa.')
        if not values['supplier_no'] or not values['internal_name']:
            raise ValueError('Numer dostawcy i nasza nazwa tworzywa są wymagane.')
        values['ordinal']=as_int(values['ordinal'],'Lp.'); values['origin_code']=as_int(values['origin_code'],'Kod oryginał/regranulat')
        values['unit_price']=as_decimal(values['unit_price']); values['price_pln']=as_decimal(values['price_pln']); values['lead_time_days']=as_int(values['lead_time_days'],'Średni czas dostawy')
        with db() as connection, connection.cursor() as cursor:
            cursor.execute('SELECT id FROM suppliers WHERE supplier_no=%s',(values['supplier_no'],)); supplier=cursor.fetchone()
        values['supplier_id']=supplier['id'] if supplier else None
        return values

    @app.get('/tworzywa')
    @auth
    def materials_index():
        q=request.args.get('q','').strip(); kind=request.args.get('kind','').strip(); color=request.args.get('color','').strip(); category=request.args.get('category','').strip()
        sql='SELECT m.*,s.name AS linked_supplier_name FROM materials m LEFT JOIN suppliers s ON s.id=m.supplier_id WHERE TRUE'; params=[]
        if q:
            term=f'%{q}%'; sql+=' AND (m.material_index ILIKE %s OR m.internal_name ILIKE %s OR m.symbol ILIKE %s OR m.supplier_no ILIKE %s OR m.supplier_name ILIKE %s)'; params += [term]*5
        if kind: sql+=' AND m.kind=%s'; params.append(kind)
        if color: sql+=' AND m.color=%s'; params.append(color)
        if category: sql+=' AND m.material_type=%s'; params.append(category)
        sql+=' ORDER BY m.material_index'
        with db() as connection, connection.cursor() as cursor:
            cursor.execute(sql,params); rows=cursor.fetchall()
            cursor.execute("SELECT DISTINCT kind FROM materials WHERE kind<>'' ORDER BY kind"); kinds=[r['kind'] for r in cursor.fetchall()]
            cursor.execute("SELECT DISTINCT color FROM materials WHERE color<>'' ORDER BY color"); colors=[r['color'] for r in cursor.fetchall()]
            cursor.execute('SELECT COUNT(*) AS count FROM materials'); total=cursor.fetchone()['count']
        return render_template('materials.html',materials=rows,kinds=kinds,colors=colors,total=total,query=q,selected_kind=kind,selected_color=color,selected_category=category)

    @app.route('/tworzywa/nowe',methods=['GET','POST'])
    @auth
    def create_material():
        if request.method=='POST':
            require_csrf()
            try:
                values=form_values()
                with db() as connection, connection.cursor() as cursor:
                    values['material_index']=next_index(cursor,index_prefix(values['material_type']))
                    columns=','.join(MATERIAL_FIELDS); placeholders=','.join('%s' for _ in MATERIAL_FIELDS)
                    cursor.execute(f'INSERT INTO materials ({columns}) VALUES ({placeholders})',[values[f] for f in MATERIAL_FIELDS])
                flash(f"Tworzywo {values['material_index']} zostało dodane.",'success'); return redirect(url_for('materials_index'))
            except ValueError as error: flash(str(error),'error')
        return render_template('material_form.html',material=request.form or {},title='Nowe tworzywo')

    @app.route('/tworzywa/<int:material_id>/edytuj',methods=['GET','POST'])
    @auth
    def edit_material(material_id):
        with db() as connection, connection.cursor() as cursor:
            cursor.execute('SELECT * FROM materials WHERE id=%s',(material_id,)); material=cursor.fetchone()
        if material is None: abort(404)
        if request.method=='POST':
            require_csrf()
            try:
                values=form_values(); prefix=index_prefix(values['material_type']); old_prefix=str(material['material_index'])[:2]
                values['material_index']=material['material_index']
                with db() as connection, connection.cursor() as cursor:
                    if old_prefix!=prefix: values['material_index']=next_index(cursor,prefix)
                    assignments=','.join(f'{f}=%s' for f in FORM_FIELDS+['supplier_id','material_index'])
                    cursor.execute(f'UPDATE materials SET {assignments},updated_at=NOW() WHERE id=%s',[values[f] for f in FORM_FIELDS+['supplier_id','material_index']]+[material_id])
                flash('Dane tworzywa zostały zapisane.','success'); return redirect(url_for('materials_index'))
            except ValueError as error: flash(str(error),'error'); material=request.form
        return render_template('material_form.html',material=material,title='Edycja tworzywa')

    @app.post('/tworzywa/<int:material_id>/usun')
    @auth
    def delete_material(material_id):
        require_csrf()
        try:
            with db() as connection, connection.cursor() as cursor: cursor.execute('DELETE FROM materials WHERE id=%s',(material_id,))
        except ForeignKeyViolation:
            flash('Tej pozycji nie można usunąć — ma historię ruchów (przyjęcia, wydania lub zwroty).','error'); return redirect(url_for('edit_material',material_id=material_id))
        flash('Tworzywo zostało usunięte.','success'); return redirect(url_for('materials_index'))
